"""Featherless Chat — a self-hosted chat UI for featherless.ai (or any OpenAI-compatible API).

Single-user. Chats, files, presets and settings live in SQLite under DATA_DIR so every
device you open the app on (phone, laptop) sees the same history.
"""

import base64
import hmac
import io
import json
import os
import secrets
import sqlite3
import time
import uuid
from functools import wraps
from pathlib import Path

import requests
from flask import (
    Flask,
    Response,
    abort,
    g,
    jsonify,
    redirect,
    request,
    send_file,
    send_from_directory,
    session,
    stream_with_context,
)

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("DATA_DIR", BASE_DIR / "data"))
UPLOAD_DIR = DATA_DIR / "uploads"
DB_PATH = DATA_DIR / "chat.db"
STATIC_DIR = BASE_DIR / "static"

DEFAULT_BASE_URL = "https://api.featherless.ai/v1"
APP_PASSWORD = os.environ.get("APP_PASSWORD", "")
MAX_UPLOAD_MB = int(os.environ.get("MAX_UPLOAD_MB", "50"))

# Sampling params forwarded to the API when set. top_k / min_p / repetition_penalty are
# not OpenAI-standard; Featherless accepts them (vLLM backend), other providers may not.
PARAM_KEYS = {
    "temperature": float,
    "top_p": float,
    "top_k": int,
    "min_p": float,
    "max_tokens": int,
    "presence_penalty": float,
    "frequency_penalty": float,
    "repetition_penalty": float,
    "seed": int,
}

TEXT_EXTS = {
    "txt", "md", "markdown", "csv", "tsv", "json", "jsonl", "xml", "yaml", "yml", "toml",
    "ini", "cfg", "conf", "log", "html", "htm", "css", "js", "mjs", "cjs", "ts", "tsx",
    "jsx", "py", "rb", "go", "rs", "java", "kt", "kts", "c", "h", "cpp", "hpp", "cc",
    "cs", "php", "swift", "m", "sh", "bash", "zsh", "ps1", "bat", "sql", "r", "lua",
    "pl", "scala", "dart", "vue", "svelte", "tex", "rst", "srt", "vtt", "env", "gitignore",
    "dockerfile", "makefile", "gradle", "properties", "ipynb",
}
IMAGE_MIMES = {"image/png", "image/jpeg", "image/webp", "image/gif"}

DATA_DIR.mkdir(parents=True, exist_ok=True)
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)


def _secret_key():
    env = os.environ.get("SECRET_KEY")
    if env:
        return env
    f = DATA_DIR / "secret_key"
    if not f.exists():
        f.write_text(secrets.token_hex(32))
    return f.read_text().strip()


app = Flask(__name__, static_folder=None)
app.config.update(
    SECRET_KEY=_secret_key(),
    MAX_CONTENT_LENGTH=MAX_UPLOAD_MB * 1024 * 1024,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE") == "1",
    PERMANENT_SESSION_LIFETIME=60 * 60 * 24 * 90,
)

# ---------------------------------------------------------------- database

SCHEMA = """
CREATE TABLE IF NOT EXISTS chats (
    id TEXT PRIMARY KEY,
    title TEXT NOT NULL DEFAULT 'New chat',
    model TEXT,
    system_prompt TEXT NOT NULL DEFAULT '',
    params TEXT NOT NULL DEFAULT '{}',
    messages TEXT NOT NULL DEFAULT '[]',
    pinned INTEGER NOT NULL DEFAULT 0,
    archived INTEGER NOT NULL DEFAULT 0,
    folder TEXT NOT NULL DEFAULT '',
    share_token TEXT UNIQUE,
    created_at REAL NOT NULL,
    updated_at REAL NOT NULL
);
CREATE TABLE IF NOT EXISTS files (
    id TEXT PRIMARY KEY,
    name TEXT NOT NULL,
    mime TEXT NOT NULL,
    size INTEGER NOT NULL,
    kind TEXT NOT NULL,
    text TEXT,
    created_at REAL NOT NULL
);
CREATE TABLE IF NOT EXISTS presets (
    id TEXT PRIMARY KEY,
    name TEXT NOT NULL,
    model TEXT,
    system_prompt TEXT NOT NULL DEFAULT '',
    params TEXT NOT NULL DEFAULT '{}',
    created_at REAL NOT NULL
);
CREATE TABLE IF NOT EXISTS settings (
    key TEXT PRIMARY KEY,
    value TEXT NOT NULL
);
"""


def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA journal_mode=WAL")
    return g.db


@app.teardown_appcontext
def _close_db(_exc):
    conn = g.pop("db", None)
    if conn is not None:
        conn.close()


with sqlite3.connect(DB_PATH) as _c:
    _c.executescript(SCHEMA)


def get_settings():
    rows = db().execute("SELECT key, value FROM settings").fetchall()
    return {r["key"]: json.loads(r["value"]) for r in rows}


def api_config():
    s = get_settings()
    key = s.get("api_key") or os.environ.get("FEATHERLESS_API_KEY", "")
    base = (s.get("base_url") or os.environ.get("FEATHERLESS_BASE_URL") or DEFAULT_BASE_URL).rstrip("/")
    return key, base


def chat_row(row, full=True):
    d = {
        "id": row["id"],
        "title": row["title"],
        "model": row["model"],
        "pinned": bool(row["pinned"]),
        "archived": bool(row["archived"]),
        "folder": row["folder"],
        "shared": bool(row["share_token"]),
        "share_token": row["share_token"],
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
    }
    if full:
        d["system_prompt"] = row["system_prompt"]
        d["params"] = json.loads(row["params"])
        d["messages"] = json.loads(row["messages"])
    return d


# ---------------------------------------------------------------- auth


def auth_required(fn):
    @wraps(fn)
    def wrapper(*a, **kw):
        if APP_PASSWORD and not session.get("ok"):
            return jsonify(error="unauthorized"), 401
        return fn(*a, **kw)

    return wrapper


@app.post("/api/login")
def login():
    pw = (request.get_json(silent=True) or {}).get("password", "")
    if not APP_PASSWORD or hmac.compare_digest(pw.encode(), APP_PASSWORD.encode()):
        session.permanent = True
        session["ok"] = True
        return jsonify(ok=True)
    time.sleep(1)  # slow down guessing
    return jsonify(error="wrong password"), 401


@app.post("/api/logout")
def logout():
    session.clear()
    return jsonify(ok=True)


@app.get("/api/me")
def me():
    return jsonify(authed=bool(not APP_PASSWORD or session.get("ok")), password_set=bool(APP_PASSWORD))


# ---------------------------------------------------------------- static / PWA


@app.get("/")
def index():
    return send_from_directory(STATIC_DIR, "index.html")


@app.get("/sw.js")
def service_worker():
    resp = send_from_directory(STATIC_DIR, "sw.js")
    resp.headers["Cache-Control"] = "no-cache"
    return resp


@app.get("/manifest.webmanifest")
def manifest():
    return send_from_directory(STATIC_DIR, "manifest.webmanifest", mimetype="application/manifest+json")


@app.get("/static/<path:p>")
def static_files(p):
    return send_from_directory(STATIC_DIR, p)


# ---------------------------------------------------------------- settings


@app.get("/api/settings")
@auth_required
def settings_get():
    s = get_settings()
    key = s.pop("api_key", "") or os.environ.get("FEATHERLESS_API_KEY", "")
    s["api_key_set"] = bool(key)
    s["api_key_hint"] = ("…" + key[-4:]) if key else ""
    s.setdefault("base_url", os.environ.get("FEATHERLESS_BASE_URL") or DEFAULT_BASE_URL)
    return jsonify(s)


@app.put("/api/settings")
@auth_required
def settings_put():
    data = request.get_json(force=True) or {}
    conn = db()
    for k, v in data.items():
        if k in ("api_key_set", "api_key_hint"):
            continue
        if k == "api_key" and v == "":
            continue  # empty field means "keep current key"
        conn.execute(
            "INSERT INTO settings(key, value) VALUES(?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (k, json.dumps(v)),
        )
    conn.commit()
    return settings_get()


# ---------------------------------------------------------------- models

_models_cache = {"at": 0.0, "base": None, "data": None}


@app.get("/api/models")
@auth_required
def models():
    key, base = api_config()
    if not key:
        return jsonify(error="No API key set. Add your Featherless key in Settings."), 400
    fresh = request.args.get("refresh") == "1"
    c = _models_cache
    if not fresh and c["data"] is not None and c["base"] == base and time.time() - c["at"] < 3600:
        return jsonify(c["data"])
    try:
        r = requests.get(f"{base}/models", headers={"Authorization": f"Bearer {key}"}, timeout=30)
    except requests.RequestException as e:
        return jsonify(error=f"Could not reach {base}: {e}"), 502
    if r.status_code != 200:
        return jsonify(error=f"Model list failed ({r.status_code}): {r.text[:500]}"), 502
    raw = r.json().get("data", [])
    out = []
    for m in raw:
        # Keep a few descriptive fields if the provider returns them.
        out.append({
            "id": m.get("id"),
            "context_length": m.get("context_length") or m.get("max_model_len"),
            "model_class": m.get("model_class"),
            "available": m.get("available_on_current_plan", True),
        })
    out.sort(key=lambda m: (m["id"] or "").lower())
    c.update(at=time.time(), base=base, data=out)
    return jsonify(out)


# ---------------------------------------------------------------- chats


@app.get("/api/chats")
@auth_required
def chats_list():
    archived = 1 if request.args.get("archived") == "1" else 0
    q = (request.args.get("q") or "").strip()
    sql = "SELECT * FROM chats WHERE archived=?"
    args = [archived]
    if q:
        sql += " AND (title LIKE ? OR messages LIKE ?)"
        args += [f"%{q}%", f"%{q}%"]
    sql += " ORDER BY pinned DESC, updated_at DESC"
    rows = db().execute(sql, args).fetchall()
    return jsonify([chat_row(r, full=False) for r in rows])


def _save_chat(cid, data, create=False):
    now = time.time()
    conn = db()
    fields = {
        "title": data.get("title"),
        "model": data.get("model"),
        "system_prompt": data.get("system_prompt"),
        "params": json.dumps(data["params"]) if "params" in data else None,
        "messages": json.dumps(data["messages"]) if "messages" in data else None,
        "pinned": int(bool(data["pinned"])) if "pinned" in data else None,
        "archived": int(bool(data["archived"])) if "archived" in data else None,
        "folder": data.get("folder"),
    }
    fields = {k: v for k, v in fields.items() if v is not None}
    if create:
        fields.setdefault("title", "New chat")
        fields.setdefault("system_prompt", "")
        cols = ["id", "created_at", "updated_at", *fields]
        conn.execute(
            f"INSERT INTO chats({','.join(cols)}) VALUES({','.join('?' * len(cols))})",
            [cid, data.get("created_at", now), data.get("updated_at", now), *fields.values()],
        )
    else:
        # Pin/archive/rename shouldn't bump the chat to the top of "recent".
        touch = "messages" in fields
        sets = ", ".join(f"{k}=?" for k in fields)
        if touch:
            sets = (sets + ", " if sets else "") + "updated_at=?"
        if not sets:
            return
        vals = [*fields.values()] + ([now] if touch else [])
        cur = conn.execute(f"UPDATE chats SET {sets} WHERE id=?", [*vals, cid])
        if cur.rowcount == 0:
            abort(404)
    conn.commit()


@app.post("/api/chats")
@auth_required
def chats_create():
    data = request.get_json(force=True) or {}
    cid = uuid.uuid4().hex
    _save_chat(cid, data, create=True)
    return chat_get(cid)


@app.get("/api/chats/<cid>")
@auth_required
def chat_get(cid):
    row = db().execute("SELECT * FROM chats WHERE id=?", (cid,)).fetchone()
    if not row:
        abort(404)
    return jsonify(chat_row(row))


@app.patch("/api/chats/<cid>")
@auth_required
def chat_patch(cid):
    _save_chat(cid, request.get_json(force=True) or {})
    return chat_get(cid)


@app.delete("/api/chats/<cid>")
@auth_required
def chat_delete(cid):
    db().execute("DELETE FROM chats WHERE id=?", (cid,))
    db().commit()
    return jsonify(ok=True)


@app.post("/api/chats/<cid>/share")
@auth_required
def chat_share(cid):
    token = secrets.token_urlsafe(16)
    cur = db().execute("UPDATE chats SET share_token=? WHERE id=?", (token, cid))
    if cur.rowcount == 0:
        abort(404)
    db().commit()
    return jsonify(token=token, url=request.host_url.rstrip("/") + f"/s/{token}")


@app.delete("/api/chats/<cid>/share")
@auth_required
def chat_unshare(cid):
    db().execute("UPDATE chats SET share_token=NULL WHERE id=?", (cid,))
    db().commit()
    return jsonify(ok=True)


def chat_markdown(chat):
    lines = [f"# {chat['title']}", ""]
    if chat.get("model"):
        lines += [f"_Model: {chat['model']}_", ""]
    if chat.get("system_prompt"):
        lines += ["**System prompt:**", "", chat["system_prompt"], ""]
    for m in chat["messages"]:
        who = "You" if m["role"] == "user" else (m.get("model") or "Assistant")
        lines += [f"## {who}", ""]
        for a in m.get("attachments") or []:
            lines.append(f"📎 {a.get('name', 'file')}")
        lines += [m.get("content", ""), ""]
    return "\n".join(lines)


@app.get("/api/chats/<cid>/export")
@auth_required
def chat_export(cid):
    row = db().execute("SELECT * FROM chats WHERE id=?", (cid,)).fetchone()
    if not row:
        abort(404)
    chat = chat_row(row)
    safe = "".join(ch for ch in chat["title"] if ch.isalnum() or ch in " -_")[:60].strip() or "chat"
    if request.args.get("format") == "md":
        return Response(chat_markdown(chat), mimetype="text/markdown",
                        headers={"Content-Disposition": f'attachment; filename="{safe}.md"'})
    chat.pop("share_token", None)
    return Response(json.dumps(chat, indent=2), mimetype="application/json",
                    headers={"Content-Disposition": f'attachment; filename="{safe}.json"'})


@app.get("/api/export")
@auth_required
def export_all():
    rows = db().execute("SELECT * FROM chats ORDER BY created_at").fetchall()
    chats = [chat_row(r) for r in rows]
    for c in chats:
        c.pop("share_token", None)
    presets = [preset_row(r) for r in db().execute("SELECT * FROM presets").fetchall()]
    s = get_settings()
    s.pop("api_key", None)
    body = json.dumps({"version": 1, "chats": chats, "presets": presets, "settings": s}, indent=2)
    return Response(body, mimetype="application/json",
                    headers={"Content-Disposition": 'attachment; filename="featherless-chat-backup.json"'})


@app.post("/api/import")
@auth_required
def import_data():
    data = request.get_json(force=True)
    chats = data if isinstance(data, list) else data.get("chats") or ([data] if "messages" in data else [])
    n = 0
    for c in chats:
        c = dict(c)
        c.pop("share_token", None)
        _save_chat(uuid.uuid4().hex, c, create=True)  # new id: never overwrite
        n += 1
    for p in (data.get("presets") or []) if isinstance(data, dict) else []:
        _save_preset(uuid.uuid4().hex, p, create=True)
    return jsonify(imported=n)


# ---------------------------------------------------------------- presets


def preset_row(r):
    return {"id": r["id"], "name": r["name"], "model": r["model"], "system_prompt": r["system_prompt"],
            "params": json.loads(r["params"]), "created_at": r["created_at"]}


def _save_preset(pid, d, create=False):
    vals = (d.get("name") or "Untitled", d.get("model"), d.get("system_prompt") or "", json.dumps(d.get("params") or {}))
    if create:
        db().execute("INSERT INTO presets(id, name, model, system_prompt, params, created_at) VALUES(?,?,?,?,?,?)",
                     (pid, *vals, time.time()))
    else:
        db().execute("UPDATE presets SET name=?, model=?, system_prompt=?, params=? WHERE id=?", (*vals, pid))
    db().commit()


@app.get("/api/presets")
@auth_required
def presets_list():
    rows = db().execute("SELECT * FROM presets ORDER BY name COLLATE NOCASE").fetchall()
    return jsonify([preset_row(r) for r in rows])


@app.post("/api/presets")
@auth_required
def presets_create():
    pid = uuid.uuid4().hex
    _save_preset(pid, request.get_json(force=True) or {}, create=True)
    return jsonify(preset_row(db().execute("SELECT * FROM presets WHERE id=?", (pid,)).fetchone()))


@app.put("/api/presets/<pid>")
@auth_required
def presets_update(pid):
    _save_preset(pid, request.get_json(force=True) or {})
    row = db().execute("SELECT * FROM presets WHERE id=?", (pid,)).fetchone()
    if not row:
        abort(404)
    return jsonify(preset_row(row))


@app.delete("/api/presets/<pid>")
@auth_required
def presets_delete(pid):
    db().execute("DELETE FROM presets WHERE id=?", (pid,))
    db().commit()
    return jsonify(ok=True)


# ---------------------------------------------------------------- files


def extract(name, mime, blob):
    """Return (kind, text). kind is 'image', 'text' or 'binary'."""
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else name.lower()
    if mime in IMAGE_MIMES:
        return "image", None
    try:
        if ext == "pdf" or mime == "application/pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(blob))
            pages = [f"[page {i + 1}]\n{p.extract_text() or ''}" for i, p in enumerate(reader.pages)]
            return "text", "\n\n".join(pages)
        if ext == "docx":
            import docx

            d = docx.Document(io.BytesIO(blob))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "text", "\n".join(parts)
    except Exception as e:  # corrupt or encrypted document
        return "binary", f"(could not extract text: {e})"
    head = blob[:8192]
    if ext in TEXT_EXTS or mime.startswith("text/") or b"\x00" not in head:
        try:
            return "text", blob.decode("utf-8")
        except UnicodeDecodeError:
            if ext in TEXT_EXTS or mime.startswith("text/"):
                return "text", blob.decode("latin-1")
    return "binary", None


@app.post("/api/files")
@auth_required
def files_upload():
    out = []
    for f in request.files.getlist("file"):
        blob = f.read()
        name = os.path.basename(f.filename or "file")
        mime = f.mimetype or "application/octet-stream"
        kind, text = extract(name, mime, blob)
        fid = uuid.uuid4().hex
        (UPLOAD_DIR / fid).write_bytes(blob)
        db().execute("INSERT INTO files(id, name, mime, size, kind, text, created_at) VALUES(?,?,?,?,?,?,?)",
                     (fid, name, mime, len(blob), kind, text, time.time()))
        out.append({"id": fid, "name": name, "mime": mime, "size": len(blob), "kind": kind,
                    "chars": len(text) if kind == "text" and text else 0})
    db().commit()
    return jsonify(out)


@app.get("/api/files/<fid>")
@auth_required
def files_get(fid):
    row = db().execute("SELECT * FROM files WHERE id=?", (fid,)).fetchone()
    if not row:
        abort(404)
    return send_file(UPLOAD_DIR / fid, mimetype=row["mime"], download_name=row["name"],
                     as_attachment=request.args.get("dl") == "1")


# ---------------------------------------------------------------- completion


def build_messages(system_prompt, messages, max_file_chars):
    out = []
    if system_prompt:
        out.append({"role": "system", "content": system_prompt})
    for m in messages:
        role = m.get("role")
        if role not in ("user", "assistant", "system"):
            continue
        text = m.get("content") or ""
        images = []
        file_blocks = []
        for a in m.get("attachments") or []:
            row = db().execute("SELECT * FROM files WHERE id=?", (a.get("id"),)).fetchone()
            if not row:
                continue
            if row["kind"] == "image":
                b64 = base64.b64encode((UPLOAD_DIR / row["id"]).read_bytes()).decode()
                images.append({"type": "image_url", "image_url": {"url": f"data:{row['mime']};base64,{b64}"}})
            elif row["kind"] == "text" and row["text"]:
                body = row["text"]
                if len(body) > max_file_chars:
                    body = body[:max_file_chars] + f"\n…[truncated, {len(row['text']) - max_file_chars} more characters]"
                file_blocks.append(f'<file name="{row["name"]}">\n{body}\n</file>')
            else:
                note = f" {row['text']}" if row["text"] else ""
                file_blocks.append(f'<file name="{row["name"]}" type="{row["mime"]}" size="{row["size"]}">'
                                   f"[binary file — contents not readable as text]{note}</file>")
        if file_blocks:
            text = "\n\n".join(file_blocks) + ("\n\n" + text if text else "")
        if images and role == "user":
            out.append({"role": role, "content": [{"type": "text", "text": text}, *images]})
        else:
            out.append({"role": role, "content": text})
    return out


def build_payload(body, settings):
    params = body.get("params") or {}
    payload = {
        "model": body.get("model"),
        "messages": build_messages(body.get("system_prompt") or "", body.get("messages") or [],
                                   int(settings.get("max_file_chars") or 200_000)),
        "stream": True,
    }
    for k, cast in PARAM_KEYS.items():
        v = params.get(k)
        if v is None or v == "":
            continue
        try:
            payload[k] = cast(v)
        except (TypeError, ValueError):
            pass
    stop = params.get("stop")
    if stop:
        payload["stop"] = [s for s in (stop if isinstance(stop, list) else str(stop).split("\n")) if s][:4]
    extra = params.get("extra")
    if extra:
        try:
            payload.update(json.loads(extra) if isinstance(extra, str) else extra)
        except json.JSONDecodeError:
            pass
    return payload


@app.post("/api/complete")
@auth_required
def complete():
    """Stream a chat completion. Body: {model, system_prompt, params, messages:[{role, content, attachments}]}.

    Response is the provider's SSE stream passed through unchanged, so the browser
    parses standard OpenAI `data: {...}` chunks.
    """
    body = request.get_json(force=True) or {}
    key, base = api_config()
    if not key:
        return jsonify(error="No API key set. Add your Featherless key in Settings."), 400
    if not body.get("model"):
        return jsonify(error="Pick a model first."), 400
    payload = build_payload(body, get_settings())
    if body.get("stream") is False:
        payload["stream"] = False
    try:
        upstream = requests.post(
            f"{base}/chat/completions",
            headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
            json=payload,
            stream=payload["stream"],
            timeout=(15, 600),
        )
    except requests.RequestException as e:
        return jsonify(error=f"Could not reach {base}: {e}"), 502
    if upstream.status_code != 200:
        text = upstream.text[:2000]
        try:
            j = upstream.json()
            err = j.get("error")
            text = (err.get("message") if isinstance(err, dict) else err) or j.get("detail") or text
        except ValueError:
            pass
        return jsonify(error=f"{upstream.status_code}: {text}"), upstream.status_code
    if not payload["stream"]:
        return jsonify(upstream.json())

    def gen():
        try:
            for chunk in upstream.iter_content(chunk_size=None):
                if chunk:
                    yield chunk
        finally:
            upstream.close()  # browser hit Stop → free the Featherless concurrency slot

    return Response(stream_with_context(gen()), mimetype="text/event-stream",
                    headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


# ---------------------------------------------------------------- public share page

SHARE_HTML = """<!doctype html><html><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title}</title><link rel="stylesheet" href="/static/style.css">
<link rel="stylesheet" href="/static/vendor/hljs-dark.css">
</head><body class="share-page"><main class="share-wrap"><h1 id="t"></h1><p class="muted" id="m"></p>
<div id="msgs" class="messages"></div><p class="muted share-foot">Shared from Featherless Chat · read-only</p></main>
<script src="/static/vendor/marked.min.js"></script><script src="/static/vendor/purify.min.js"></script>
<script src="/static/vendor/highlight.min.js"></script>
<script>const CHAT={data};</script><script src="/static/share.js"></script></body></html>"""


@app.get("/s/<token>")
def shared(token):
    row = db().execute("SELECT * FROM chats WHERE share_token=?", (token,)).fetchone()
    if not row:
        abort(404)
    c = chat_row(row)
    data = {"title": c["title"], "model": c["model"],
            "messages": [{"role": m["role"], "content": m.get("content", ""), "model": m.get("model"),
                          "attachments": [{"name": a.get("name")} for a in m.get("attachments") or []]}
                         for m in c["messages"]]}
    # json.dumps output is safe inside <script> once '</' is escaped.
    blob = json.dumps(data).replace("</", "<\\/")
    title = "".join(ch for ch in c["title"] if ch not in "<>&\"'{}")
    return SHARE_HTML.replace("{title}", title).replace("{data}", blob)


if __name__ == "__main__":
    if not APP_PASSWORD:
        print("WARNING: APP_PASSWORD is not set — anyone who can reach this server can use it.")
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, threaded=True)
